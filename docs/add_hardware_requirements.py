"""Add the client-supplied installation requirements without replacing existing guide edits."""

from pathlib import Path
from docx import Document

SOURCE = Path('docs/TDQSYS_CLIENT_HANDOVER_GUIDE.docx')
OUTPUT = Path('docs/TDQSYS_CLIENT_HANDOVER_GUIDE_FINAL.docx')
ANCHOR = 'TDQSYS is a local queue management system for vehicle test drives.'


def add_after(anchor, paragraph):
    anchor._p.addnext(paragraph._p)
    return paragraph


doc = Document(SOURCE)
anchor = next((p for p in doc.paragraphs if p.text.startswith(ANCHOR)), None)
if anchor is None:
    raise RuntimeError('Could not find the guide introduction to place the hardware requirements.')

cursor = anchor
heading = doc.add_heading('Minimum Hardware and Installation Requirements', level=1)
cursor = add_after(cursor, heading)

intro = doc.add_paragraph(
    'Prepare the following equipment before installing TDQSYS. Place the computer where the operator can use it comfortably and where the display cable can be routed safely.'
)
cursor = add_after(cursor, intro)

for item in [
    'Windows PC or laptop to run TDQSYS.',
    'External TV or monitor of at least 43 inches for the customer-facing queue display.',
    'HDMI cable long enough for the measured distance between the computer and the external display, with enough allowance for safe routing and connection.',
]:
    bullet = doc.add_paragraph(item, style='List Bullet')
    cursor = add_after(cursor, bullet)

doc.save(OUTPUT)
