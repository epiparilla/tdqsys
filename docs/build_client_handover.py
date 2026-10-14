import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = sys.argv[1] if len(sys.argv) > 1 else 'docs/TDQSYS_CLIENT_HANDOVER_GUIDE.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)

for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
    s = styles[name]
    s.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    s._element.rPr.rFonts.set(qn('w:ascii'), s.font.name)
    s._element.rPr.rFonts.set(qn('w:hAnsi'), s.font.name)
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.bold = name != 'Title'
    s.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(7)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)

def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c, '17365D')
        for r in c.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(255,255,255)
            r.font.bold = True
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(c)
        set_cell_margin(c)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            if ri % 2 == 1:
                shade(cells[i], 'F3F6F9')
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[i])
            set_cell_margin(cells[i])
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return t

def heading(text, level=1):
    return doc.add_heading(text, level=level)

def p(text='', bold_lead=None):
    par = doc.add_paragraph()
    if bold_lead:
        run = par.add_run(bold_lead)
        run.bold = True
        par.add_run(text)
    else:
        par.add_run(text)
    return par

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

# Title page
title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('TDQSYS Client Handover Guide')
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Queue management, TV display, licensing, backup and support')
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(70, 70, 70)
doc.add_paragraph()
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.add_run('Version 1.1.27').bold = True
doc.add_paragraph()
p('This guide is for the site operator and manager who use TDQSYS day to day. It explains how to operate the queue, manage the TV display, keep a safe backup, activate or renew a license, and obtain support. Keep this document with the site handover records.')
doc.add_page_break()

heading('What TDQSYS Does')
p('TDQSYS is a local queue management system for vehicle test drives. The desk operator changes queue numbers from the dashboard. The TV display shows the live queue, plays optional video adverts, and announces queue calls. The system continues to operate locally when the internet is unavailable.')
table(['Part', 'Purpose'], [
    ('Facilitator dashboard', 'Used by the desk operator to increase, decrease, type a queue number, or re-announce a vehicle.'),
    ('TV display', 'Shows the queue to visitors, plays optional adverts, and performs voice announcements.'),
    ('Settings', 'Used on the TDQSYS PC for setup, licensing, updates, backups, videos, and configuration.'),
    ('Phone mirror link', 'A read-only web view for the specific location. It does not allow queue changes.'),
], [2.0, 4.5])

heading('Minimum Hardware and Installation Requirements')
p('Prepare the following equipment before installing TDQSYS. Place the computer where the operator can use it comfortably and where the display cable can be routed safely.')
bullets([
    'Windows PC or laptop to run TDQSYS.',
    'External TV or monitor of at least 43 inches for the customer-facing queue display.',
    'HDMI cable long enough for the measured distance between the computer and the external display, with enough allowance for safe routing and connection.',
])

heading('Before the Event')
bullets([
    'Start TDQSYS from its Windows shortcut and allow the dashboard to load.',
    'Confirm the TV display is visible on the external screen. TDQSYS automatically prefers a connected second display; without one it uses a local preview window.',
    'Open Settings on the TDQSYS PC. Confirm the location label, vehicle list, video count, and license status.',
    'Make a backup before changing configuration or starting a major event.',
    'If phones will be used to watch the queue, copy or print the mirror-link QR code from Settings.'
])

heading('Daily Operation')
heading('Run the Queue', 2)
p('Use the Facilitator Dashboard on the TDQSYS PC. Each vehicle has a minus button, a number field, a plus button, and a re-announce button. Queue numbers can be set from 0 to 999. Changes are saved automatically and are reflected on the TV display.')
bullets([
    'Use plus or minus for normal calls.',
    'Type a number when a larger adjustment is required, then confirm the field.',
    'Use the announcement button to repeat the current call without changing the number.',
    'Do not use Reset All during an event unless you have confirmed that all counters should return to zero.'
])
heading('Use the TV Display', 2)
p('The TV display is intended for visitors. It shows the configured brands and vehicles, their current numbers, optional video adverts, and voice announcements. The display remains viewable even when a license has expired or is temporarily revoked; only changes are disabled.')
bullets([
    'Use Open Display Screen from the dashboard if the TV window was closed.',
    'If the HDMI display is disconnected, TDQSYS changes the display to a smaller preview window. Reconnect the display and TDQSYS will detect it.',
    'Advert playback can be enabled or disabled in Settings. Imported MP4 or WebM videos are used only on the local TV display.'
])
heading('Share the Read Only Mirror', 2)
p('Settings provides an immutable Instance ID, a mirror link, and a QR code. Share the mirror link or QR code with staff or guests who need to view this location remotely. The mirror is read-only and cannot change queue numbers or settings.')

heading('Settings and Configuration')
p('Settings is available only on the TDQSYS PC. Select Enter Edit Mode before changing configuration. Save Configuration when finished, or cancel to discard the edits.')
table(['Setting', 'Use'], [
    ('Location label', 'A friendly local name. It is not the system identity and duplicate labels do not mix data.'),
    ('Brands and cars', 'Configure up to two brands and a maximum of 12 vehicles in total.'),
    ('Advertising', 'Enable or disable TV adverts and set the minimum and maximum interval between adverts.'),
    ('Videos', 'Import MP4 or WebM videos from a folder or USB drive. Clear Videos removes the imported TV adverts.'),
    ('Cloud mirror', 'Leave enabled for read-only phone viewing. The local TDQSYS PC remains the source of truth.'),
], [2.0, 4.5])
p('Important: the Instance ID is the permanent identity of this installation. Do not change it. When requesting support, activation, or a license renewal, provide this ID exactly as shown in Settings.')

heading('Licensing')
p('A valid license permits queue changes, configuration saves, video import and deletion. The status appears in Settings.')
table(['Status', 'What continues to work', 'What is blocked'], [
    ('Valid', 'All features.', 'Nothing.'),
    ('Expired or not activated', 'TV display, queue viewing, phone mirror, existing data and videos.', 'Queue changes, configuration saves, and video management.'),
    ('Revoked', 'TV display, queue viewing, phone mirror, existing data and videos.', 'Queue changes, configuration saves, and video management until the provider cancels the revocation.'),
], [1.5, 2.7, 2.3])
heading('Activate or Renew', 2)
bullets([
    'Open Settings and expand Activate or renew a licence.',
    'Use Copy activation request, then send the request to your TDQSYS provider.',
    'Paste the issued code or import the supplied license text file, then select Activate.',
    'Confirm that Settings shows Valid before returning to normal operation.'
])
heading('Revocation and Reinstatement', 2)
p('Revocation is a reversible operational lock managed by the provider. It does not erase queue data, videos, or the stored activation. When the provider cancels the revocation, TDQSYS restores normal access automatically after cloud synchronization. If it does not change within about 20 seconds, restart TDQSYS once.')

heading('Software Updates')
p('Updates are offered from the dashboard and Settings. Select Check for Updates, then Download and Install Update. TDQSYS downloads the verified installer, installs it, and restarts. An internet connection and a Windows administrator approval prompt may be required.')
p('For a new installation, share this permanent client download link: https://tdqsys.pages.dev/download. It always finds the current Windows installer, even when a later release uses a different file name.', bold_lead='Latest installer: ')
bullets([
    'Do not power off the PC while the update is downloading or installing.',
    'After restart, confirm the version shown in Settings.',
    'Updates preserve local queue data, videos, and a valid activation.'
])

heading('Backup and Restore')
p('Backup creates one ZIP file containing the queue state, configuration, and imported videos. Keep the ZIP in a secure location separate from the TDQSYS PC, such as approved cloud storage or an encrypted USB drive.')
table(['When', 'Action'], [
    ('Before an event', 'Create a backup after configuration is confirmed.'),
    ('After major configuration changes', 'Create another backup.'),
    ('Before moving or uninstalling TDQSYS', 'Create and verify a backup ZIP.'),
    ('After replacement PC setup', 'Restore the backup from Settings, then confirm the location and queue data.'),
], [2.5, 4.0])
p('A data backup does not replace the separate license activation process. If TDQSYS is moved to a new PC, contact the provider with the new Instance ID for the correct activation procedure.')

heading('Uninstall and Replacement PC')
p('Uninstall is available from Settings on the local TDQSYS PC. It is intended to leave the machine clean. Create a backup first. After uninstalling, a new installation or replacement PC requires a restore of the backup and confirmation of the licensing arrangement with the provider.')

heading('Quick Troubleshooting')
table(['Situation', 'Check first'], [
    ('TV display is not visible', 'Reconnect or power on the external display, then use Open Display Screen from the dashboard.'),
    ('A queue change does not save', 'Check the license banner in the dashboard or Settings. Confirm the PC has not lost local storage space.'),
    ('Phone mirror does not update', 'Check the TDQSYS PC is running and online. Reopen the exact mirror link from Settings.'),
    ('No voice announcement or adverts', 'Confirm the TV display is open, adverts are enabled if required, and videos appear in Settings.'),
    ('License shows expired, revoked, or not activated', 'Open Settings, copy the activation request and contact the provider. Do not uninstall the app or reset the PC.'),
    ('Need to recover after a PC problem', 'Install TDQSYS, restore the latest backup, and contact the provider for activation if required.'),
], [2.4, 4.1])

heading('Support Information')
p('When requesting support, provide the location label, Instance ID, installed TDQSYS version, a clear description of the issue, and a photo or screenshot of any message shown. Do not send license codes to unauthorized people and do not share the activation file or Windows account credentials.')
p('For operational continuity, retain the latest backup ZIP and this handover guide with the site records.')

# Footer
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('TDQSYS Client Handover Guide  |  Version 1.1.27')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUT)
